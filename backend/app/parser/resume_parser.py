"""
Resume Parser - Extracts structured information from PDF and DOCX resumes
Uses pdfplumber for PDF, python-docx for DOCX
"""
import re
import os
from typing import Dict, List, Any, Optional, Tuple
from pathlib import Path
import pdfplumber
import docx
from loguru import logger
from app.config import settings


# Common skill patterns
SKILL_PATTERNS = [
    # Programming languages
    r'\b(Python|Java|JavaScript|TypeScript|C\+\+|C#|Go|Rust|PHP|Ruby|Swift|Kotlin|R|Scala|MATLAB)\b',
    # Web frameworks
    r'\b(React|Angular|Vue\.?js|Next\.?js|Django|FastAPI|Flask|Spring|Express|Laravel|Rails)\b',
    # Databases
    r'\b(MySQL|PostgreSQL|MongoDB|Redis|SQLite|Oracle|Cassandra|DynamoDB|Firebase|ElasticSearch)\b',
    # Cloud & DevOps
    r'\b(AWS|Azure|GCP|Docker|Kubernetes|Jenkins|CI/CD|Terraform|Ansible|Linux|Git|GitHub)\b',
    # AI/ML
    r'\b(Machine Learning|Deep Learning|TensorFlow|PyTorch|scikit-learn|NLP|Computer Vision|LLM|RAG)\b',
    # Other tech
    r'\b(REST API|GraphQL|Microservices|Agile|Scrum|HTML|CSS|Bootstrap|Tailwind|Node\.?js)\b',
]


class ResumeParser:
    """Parses resumes and extracts structured information"""

    def parse(self, file_path: str) -> Dict[str, Any]:
        """Main parsing entry point - detects file type and routes to parser"""
        path = Path(file_path)
        extension = path.suffix.lower()

        logger.info(f"📄 Parsing resume: {path.name}")

        if extension == ".pdf":
            full_text = self._extract_pdf_text(file_path)
        elif extension in [".docx", ".doc"]:
            full_text = self._extract_docx_text(file_path)
        else:
            raise ValueError(f"Unsupported file format: {extension}")

        if not full_text or not full_text.strip():
            raise ValueError("Could not extract text from resume. File may be image-based or corrupted.")

        # Extract structured data
        parsed = self._extract_structured_data(full_text)
        parsed["full_text"] = full_text

        logger.info(f"✅ Resume parsed: {len(full_text)} chars, {len(parsed.get('skills', []))} skills found")
        return parsed

    def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text from PDF using pdfplumber"""
        text_parts = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
        except Exception as e:
            logger.error(f"❌ PDF extraction failed: {e}")
            raise ValueError(f"Could not read PDF: {e}")

        return "\n\n".join(text_parts)

    def _extract_docx_text(self, file_path: str) -> str:
        """Extract text from DOCX"""
        try:
            doc = docx.Document(file_path)
            paragraphs = []

            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())

            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())

            return "\n".join(paragraphs)
        except Exception as e:
            logger.error(f"❌ DOCX extraction failed: {e}")
            raise ValueError(f"Could not read DOCX: {e}")

    def _extract_structured_data(self, text: str) -> Dict[str, Any]:
        """Extract structured fields from resume text"""
        return {
            "name": self._extract_name(text),
            "email": self._extract_email(text),
            "phone": self._extract_phone(text),
            "skills": self._extract_skills(text),
            "experience": self._extract_experience(text),
            "education": self._extract_education(text),
            "projects": self._extract_projects(text),
            "certifications": self._extract_certifications(text),
            "summary": self._extract_summary(text),
        }

    def _extract_name(self, text: str) -> str:
        """Extract candidate name (usually first non-empty line)"""
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        if lines:
            # First line is usually the name
            first_line = lines[0]
            # Filter out lines that look like headers or URLs
            if len(first_line.split()) <= 5 and "@" not in first_line and "http" not in first_line:
                return first_line
        return ""

    def _extract_email(self, text: str) -> str:
        """Extract email address"""
        pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        matches = re.findall(pattern, text)
        return matches[0] if matches else ""

    def _extract_phone(self, text: str) -> str:
        """Extract phone number"""
        patterns = [
            r'\+?[\d\s\-\(\)]{10,15}',
            r'\b\d{10}\b',
            r'\b\d{3}[\-\s]\d{3}[\-\s]\d{4}\b',
        ]
        for pattern in patterns:
            matches = re.findall(pattern, text)
            if matches:
                # Clean the match
                phone = re.sub(r'[^\d+\-\s]', '', matches[0]).strip()
                if len(re.sub(r'[^\d]', '', phone)) >= 10:
                    return phone
        return ""

    def _extract_skills(self, text: str) -> List[str]:
        """Extract technical skills using pattern matching + section parsing"""
        skills: set = set()

        # Step 1: Use curated regex patterns — always clean
        for pattern in SKILL_PATTERNS:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for m in matches:
                skills.add(m.strip())

        # Step 2: Parse skills section for unlisted skills
        skills_section = self._extract_section(text, ["skills", "technical skills", "technologies", "tech stack"])
        if skills_section:
            for raw_skill in re.split(r'[,|•\n\t/]', skills_section):
                skill = raw_skill.strip().strip('-').strip('•').strip()
                # Strict filters:
                if not skill or len(skill) < 2:
                    continue
                # Reject if too long (sentences)
                if len(skill) > 40:
                    continue
                # Reject if too many words (likely a sentence fragment)
                words = skill.split()
                if len(words) > 4:
                    continue
                # Reject if it starts with a verb or dash-sentence indicator
                if skill.startswith(('-', '–', '*', 'Worked', 'Supported', 'Configured',
                                     'Performed', 'Developed', 'Managed', 'Led', 'Built',
                                     'Created', 'Implemented', 'Designed', 'Maintained')):
                    continue
                # Reject if it's a digit-only string or location/university name patterns
                if skill.isdigit():
                    continue
                # Reject obvious noise (cities, months, years)
                if re.match(r'^(KOCHI|PALAKKAD|APRIL|MAY|JUNE|JULY|AUGUST|20\d\d|19\d\d)$', skill, re.IGNORECASE):
                    continue
                skills.add(skill)

        # Deduplicate case-insensitively — keep the most title-cased version
        seen_lower: dict = {}
        for s in sorted(skills):
            key = s.lower().strip()
            if key not in seen_lower:
                seen_lower[key] = s
            else:
                # Prefer the version that matches title/original case
                existing = seen_lower[key]
                if s[0].isupper() and not existing[0].isupper():
                    seen_lower[key] = s

        return list(seen_lower.values())[:50]  # Limit to 50 skills

    def _extract_experience(self, text: str) -> List[str]:
        """Extract work experience entries"""
        section = self._extract_section(
            text,
            ["experience", "work experience", "employment", "work history", "professional experience"]
        )
        if not section:
            return []

        # Split into individual experience entries
        entries = []
        current_entry = []

        for line in section.split("\n"):
            line = line.strip()
            if not line:
                if current_entry:
                    entries.append(" ".join(current_entry))
                    current_entry = []
            else:
                current_entry.append(line)

        if current_entry:
            entries.append(" ".join(current_entry))

        return entries[:10]

    def _extract_education(self, text: str) -> List[str]:
        """Extract education entries"""
        section = self._extract_section(
            text,
            ["education", "academic background", "qualifications", "academics"]
        )
        if not section:
            return []

        entries = [line.strip() for line in section.split("\n") if line.strip()]
        return entries[:10]

    def _extract_projects(self, text: str) -> List[str]:
        """Extract project descriptions"""
        section = self._extract_section(
            text,
            ["projects", "personal projects", "key projects", "academic projects"]
        )
        if not section:
            return []

        entries = []
        current_entry = []

        for line in section.split("\n"):
            line = line.strip()
            if not line:
                if current_entry:
                    entries.append(" ".join(current_entry))
                    current_entry = []
            else:
                current_entry.append(line)

        if current_entry:
            entries.append(" ".join(current_entry))

        return entries[:10]

    def _extract_certifications(self, text: str) -> List[str]:
        """Extract certifications"""
        section = self._extract_section(
            text,
            ["certifications", "certificates", "awards", "achievements", "licenses"]
        )
        if not section:
            return []

        entries = [line.strip() for line in section.split("\n") if line.strip()]
        return entries[:10]

    def _extract_summary(self, text: str) -> str:
        """Extract professional summary/objective"""
        section = self._extract_section(
            text,
            ["summary", "objective", "profile", "about me", "professional summary"]
        )
        if section:
            return section[:500]

        # If no section found, use first substantive paragraph
        paragraphs = [p.strip() for p in text.split("\n\n") if len(p.strip()) > 100]
        if len(paragraphs) > 1:
            return paragraphs[1][:500]  # Skip name/contact line
        return ""

    def _extract_section(self, text: str, section_names: List[str]) -> str:
        """Extract text under a named section header"""
        lines = text.split("\n")
        section_content = []
        in_section = False

        # Section headers for stopping
        all_section_headers = [
            "experience", "education", "skills", "projects", "certifications",
            "awards", "references", "languages", "hobbies", "interests",
            "summary", "objective", "profile", "publications", "volunteer"
        ]

        for i, line in enumerate(lines):
            line_lower = line.strip().lower()

            # Check if we're entering the target section
            if any(name in line_lower for name in section_names):
                in_section = True
                continue

            # Check if we're entering a different section (stop collecting)
            if in_section and any(
                header in line_lower
                for header in all_section_headers
                if not any(name in line_lower for name in section_names)
            ):
                # Only stop if this looks like a section header (short line)
                if len(line.strip()) < 50 and line.strip():
                    break

            if in_section:
                section_content.append(line)

        return "\n".join(section_content[:30]).strip()  # Limit to 30 lines per section

    def chunk_resume(self, parsed_data: Dict[str, Any]) -> List[str]:
        """
        Split resume into logical chunks for vector storage.
        Each section becomes a chunk for granular retrieval.
        """
        chunks = []
        chunk_size = settings.chunk_size

        # Chunk 1: Name + Summary + Skills (identity chunk)
        identity_parts = []
        if parsed_data.get("name"):
            identity_parts.append(f"Name: {parsed_data['name']}")
        if parsed_data.get("summary"):
            identity_parts.append(f"Summary: {parsed_data['summary']}")
        if parsed_data.get("skills"):
            skills_text = ", ".join(parsed_data["skills"][:30])
            identity_parts.append(f"Skills: {skills_text}")

        if identity_parts:
            chunks.append("\n".join(identity_parts))

        # Chunk 2: Education
        if parsed_data.get("education"):
            edu_text = "Education:\n" + "\n".join(
                str(e) for e in parsed_data["education"]
            )
            chunks.append(edu_text)

        # Chunk 3: Experience (split into multiple chunks if large)
        if parsed_data.get("experience"):
            exp_text = "Work Experience:\n" + "\n\n".join(
                str(e) for e in parsed_data["experience"]
            )
            # Split if too long
            if len(exp_text) > chunk_size:
                for i in range(0, len(exp_text), chunk_size):
                    chunk = exp_text[i:i + chunk_size]
                    if chunk.strip():
                        chunks.append(chunk)
            else:
                chunks.append(exp_text)

        # Chunk 4: Projects (split if needed)
        if parsed_data.get("projects"):
            proj_text = "Projects:\n" + "\n\n".join(
                str(p) for p in parsed_data["projects"]
            )
            if len(proj_text) > chunk_size:
                for i in range(0, len(proj_text), chunk_size):
                    chunk = proj_text[i:i + chunk_size]
                    if chunk.strip():
                        chunks.append(chunk)
            else:
                chunks.append(proj_text)

        # Chunk 5: Certifications
        if parsed_data.get("certifications"):
            cert_text = "Certifications:\n" + "\n".join(
                str(c) for c in parsed_data["certifications"]
            )
            chunks.append(cert_text)

        # Fallback: Use full text if no chunks extracted
        if not chunks and parsed_data.get("full_text"):
            full_text = parsed_data["full_text"]
            for i in range(0, len(full_text), chunk_size):
                chunk = full_text[i:i + chunk_size]
                if chunk.strip():
                    chunks.append(chunk)

        logger.info(f"📦 Created {len(chunks)} chunks from resume")
        return chunks


# Singleton instance
resume_parser = ResumeParser()
